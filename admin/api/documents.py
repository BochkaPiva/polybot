from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, BackgroundTasks
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select
from typing import List
import os
import shutil
import uuid
from datetime import datetime
import PyPDF2
import docx
import pytesseract
from PIL import Image
import io
import aiofiles
import mimetypes
from ..core import security
from ..core.database import get_async_session
from ..core.schemas import DocumentResponse, DocumentCreate, DocumentUpdate
from ..core.models import Document, User
from ..core.config import settings

router = APIRouter()

# Создаем директорию для загрузки файлов, если она не существует
UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

async def extract_text_from_file(file_path: str, file_type: str) -> str:
    """Извлекает текст из файла в зависимости от его типа"""
    try:
        if file_type == "application/pdf":
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                text = ""
                for page in reader.pages:
                    text += page.extract_text() + "\n"
                return text
                
        elif file_type in ["application/msword", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"]:
            doc = docx.Document(file_path)
            return "\n".join([paragraph.text for paragraph in doc.paragraphs])
            
        elif file_type.startswith("image/"):
            image = Image.open(file_path)
            return pytesseract.image_to_string(image, lang='rus+eng')
            
        elif file_type == "text/plain":
            async with aiofiles.open(file_path, mode='r', encoding='utf-8') as file:
                return await file.read()
                
        return ""
    except Exception as e:
        print(f"Error extracting text: {str(e)}")
        return ""

async def process_document(document_id: int, db: AsyncSession):
    """Фоновая обработка документа: извлечение текста"""
    try:
        # Получаем документ
        query = select(Document).where(Document.id == document_id)
        result = await db.execute(query)
        document = result.scalar_one_or_none()
        
        if not document:
            return
            
        # Получаем путь к файлу
        file_path = os.path.join(UPLOAD_DIR, document.system_filename)
        
        # Извлекаем текст
        content_text = await extract_text_from_file(file_path, document.file_type)
        
        # Обновляем документ
        document.content_text = content_text
        await db.commit()
        
    except Exception as e:
        print(f"Error processing document: {str(e)}")

@router.get("/", response_model=List[DocumentResponse])
async def get_documents(
    skip: int = 0,
    limit: int = 100,
    db: AsyncSession = Depends(get_async_session)
    # Temporarily disabled authentication
    # current_user: User = Depends(security.get_current_active_admin)
):
    query = select(Document).offset(skip).limit(limit)
    result = await db.execute(query)
    documents = result.scalars().all()
    return documents

@router.get("/{doc_id}", response_model=DocumentResponse)
async def get_document(
    doc_id: int,
    db: AsyncSession = Depends(get_async_session)
    # Temporarily disabled authentication
    # current_user: User = Depends(security.get_current_active_admin)
):
    query = select(Document).where(Document.id == doc_id)
    result = await db.execute(query)
    document = result.scalar_one_or_none()
    if document is None:
        raise HTTPException(status_code=404, detail="Document not found")
    return document

@router.post("/", response_model=DocumentResponse)
async def create_document(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_async_session)
    # Temporarily disabled authentication
    # current_user: User = Depends(security.get_current_active_admin)
):
    # Генерируем уникальное имя файла
    file_extension = os.path.splitext(file.filename)[1]
    system_filename = f"{uuid.uuid4()}{file_extension}"
    file_path = os.path.join(UPLOAD_DIR, system_filename)
    
    # Сохраняем файл
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
    
    # Определяем MIME-тип файла
    file_type, _ = mimetypes.guess_type(file.filename)
    if not file_type:
        file_type = "application/octet-stream"
    
    # Создаем запись в базе данных
    document = Document(
        original_filename=file.filename,
        system_filename=system_filename,
        file_type=file_type,
        file_size=os.path.getsize(file_path),
        created_at=datetime.utcnow()
    )
    
    db.add(document)
    await db.commit()
    await db.refresh(document)
    
    # Добавляем задачу на извлечение текста
    background_tasks.add_task(process_document, document.id, db)
    
    return document

@router.put("/{doc_id}", response_model=DocumentResponse)
async def update_document(
    doc_id: int,
    doc_update: DocumentUpdate,
    db: AsyncSession = Depends(get_async_session)
    # Temporarily disabled authentication
    # current_user: User = Depends(security.get_current_active_admin)
):
    query = select(Document).where(Document.id == doc_id)
    result = await db.execute(query)
    document = result.scalar_one_or_none()
    if document is None:
        raise HTTPException(status_code=404, detail="Document not found")
    
    update_data = doc_update.dict(exclude_unset=True)
    for field, value in update_data.items():
        setattr(document, field, value)
    
    document.updated_at = datetime.utcnow()
    await db.commit()
    await db.refresh(document)
    return document

@router.delete("/{doc_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_document(
    doc_id: int,
    db: AsyncSession = Depends(get_async_session)
    # Temporarily disabled authentication
    # current_user: User = Depends(security.get_current_active_admin)
):
    query = select(Document).where(Document.id == doc_id)
    result = await db.execute(query)
    document = result.scalar_one_or_none()
    if document is None:
        raise HTTPException(status_code=404, detail="Document not found")
    
    # Удаляем файл
    file_path = os.path.join(UPLOAD_DIR, document.system_filename)
    if os.path.exists(file_path):
        os.remove(file_path)
    
    await db.delete(document)
    await db.commit()
    return None 